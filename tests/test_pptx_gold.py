"""教会常见 PPT 金样：标题 / 经文 / 要点导入回归。"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from core.sermon.import_report import import_dialog_body, import_notes_list
from core.sermon.store import SermonStore


def _build_church_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    prs = Presentation()
    prs.slide_width = Emu(int(13.333333 * 914400))
    prs.slide_height = Emu(int(7.5 * 914400))
    blank = prs.slide_layouts[6]
    prs.core_properties.title = "主日讲道金样"

    # 1 标题页
    s1 = prs.slides.add_slide(blank)
    box = s1.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(1.2))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "恩典之路"
    run.font.size = Pt(44)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)
    sub = s1.shapes.add_textbox(Inches(2), Inches(4), Inches(9), Inches(0.6))
    sp = sub.text_frame.paragraphs[0]
    sp.alignment = PP_ALIGN.CENTER
    sr = sp.add_run()
    sr.text = "约翰福音 1:14"
    sr.font.size = Pt(24)
    sr.font.color.rgb = RGBColor(200, 200, 200)

    # 2 经文页
    s2 = prs.slides.add_slide(blank)
    ref = s2.shapes.add_textbox(Inches(1), Inches(1), Inches(11), Inches(0.6))
    rp = ref.text_frame.paragraphs[0]
    rp.alignment = PP_ALIGN.CENTER
    rr = rp.add_run()
    rr.text = "约翰福音 3:16"
    rr.font.size = Pt(22)
    rr.font.bold = True
    rr.font.color.rgb = RGBColor(255, 215, 0)
    verse = s2.shapes.add_textbox(Inches(1.2), Inches(2), Inches(10.5), Inches(3.5))
    vp = verse.text_frame.paragraphs[0]
    vp.alignment = PP_ALIGN.CENTER
    vr = vp.add_run()
    vr.text = "神爱世人，甚至将他的独生子赐给他们，叫一切信他的，不至灭亡，反得永生。"
    vr.font.size = Pt(28)
    vr.font.color.rgb = RGBColor(255, 255, 255)

    # 3 要点页
    s3 = prs.slides.add_slide(blank)
    head = s3.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11), Inches(0.8))
    hp = head.text_frame.paragraphs[0]
    hr = hp.add_run()
    hr.text = "三点提醒"
    hr.font.size = Pt(36)
    hr.font.bold = True
    hr.font.color.rgb = RGBColor(255, 255, 255)
    body = s3.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10), Inches(4))
    tf = body.text_frame
    for i, line in enumerate(["仰望十架", "彼此相爱", "广传福音"]):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = para.add_run()
        run.text = f"• {line}"
        run.font.size = Pt(28)
        run.font.color.rgb = RGBColor(240, 240, 240)

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))


class PptxGoldTests(unittest.TestCase):
    def test_church_title_scripture_bullets(self):
        try:
            from core.sermon.pptx_io import import_pptx
        except Exception:
            self.skipTest("python-pptx unavailable")
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            pptx = root / "church_gold.pptx"
            _build_church_pptx(pptx)
            store = SermonStore(root / "work")
            doc = import_pptx(pptx, store)
            self.assertEqual(doc.title, "主日讲道金样")
            self.assertEqual(len(doc.slides), 3)

            texts = [
                " ".join(el.content for el in slide.elements if el.type == "text")
                for slide in doc.slides
            ]
            self.assertIn("恩典之路", texts[0])
            self.assertIn("约翰福音 3:16", texts[1])
            self.assertIn("神爱世人", texts[1])
            self.assertIn("仰望十架", texts[2])
            self.assertIn("广传福音", texts[2])

            # 标题页字号应明显大于副标题（画布像素）
            title_els = [e for e in doc.slides[0].elements if e.type == "text"]
            self.assertGreaterEqual(len(title_els), 2)
            sizes = sorted((e.style.font_size for e in title_els), reverse=True)
            self.assertGreater(sizes[0], sizes[1])

    def test_import_report_helpers(self):
        class _Doc:
            import_skip_count = 3
            import_notes = ["第1页：跳过图表", "第2页：跳过表格（请用 Office 编辑）"]

        notes = import_notes_list(_Doc())
        self.assertTrue(any("图表" in n for n in notes))
        body = import_dialog_body(_Doc())
        self.assertIsNotNone(body)
        self.assertIn("图表", body or "")


if __name__ == "__main__":
    unittest.main()
