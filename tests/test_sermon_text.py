from __future__ import annotations

import sys
import unittest
from pathlib import Path

from PyQt6.QtGui import QFont, QTextCursor
from PyQt6.QtWidgets import QApplication, QGraphicsTextItem

from core.sermon.model import Background, ElementStyle, new_slide, new_text_element
from core.sermon.pptx_io import (
    _expand_pptx_tabs,
    _font_name_implies_bold,
    _html_escape,
    _iter_shapes,
    _shape_text_content,
)
from ui.sermon.text_format import (
    RICH_HTML_MARK,
    RICH_RUNS_MARK,
    apply_text_content,
    encode_rich_runs,
    is_text_placeholder,
    parse_rich_runs,
    plain_from_content,
    resolve_qfont,
    weight_from_font_name,
)

_APP = QApplication.instance() or QApplication(sys.argv)


class SermonTextTests(unittest.TestCase):
    def test_placeholder(self):
        self.assertTrue(is_text_placeholder("在此输入"))
        self.assertTrue(is_text_placeholder("  双击编辑文字  "))
        self.assertFalse(is_text_placeholder("约翰福音"))
        self.assertFalse(is_text_placeholder(""))
        runs = encode_rich_runs([("在此输入", "#FFFFFF")])
        self.assertTrue(is_text_placeholder(runs))

    def test_new_text_copies_style(self):
        style = ElementStyle(font_size=36, color="#FFAA00", line_spacing=160)
        el = new_text_element("在此输入", style=style)
        self.assertEqual(el.style.font_size, 36)
        self.assertEqual(el.style.color, "#FFAA00")
        style.font_size = 12
        self.assertEqual(el.style.font_size, 36)

    def test_new_slide_clones_background(self):
        bg = Background(type="color", value="#112233")
        slide = new_slide(background=bg)
        self.assertEqual(slide.background.value, "#112233")
        self.assertIsNot(slide.background, bg)

    def test_html_keeps_spaces_and_tabs(self):
        html = _html_escape("1.成就 我做了什么")
        self.assertIn("&nbsp;", html)
        html2 = _html_escape("复\u3000活")
        self.assertIn("&#12288;", html2)

    def test_tabs_kept_for_tab_stops(self):
        self.assertEqual(_expand_pptx_tabs("复\t活"), "复\t活")

    def test_missing_font_falls_back_to_installed(self):
        from PyQt6.QtGui import QFontDatabase

        installed = {f.lower() for f in QFontDatabase.families()}
        if not installed:
            self.skipTest("QFontDatabase empty (no GUI font stack)")
        missing = "ZZNoSuchFontFamilyXYZ 95 ExtraBold"
        self.assertFalse(any(missing.split()[0].lower() in f for f in installed))
        font = resolve_qfont(missing, pixel_size=48, bold=True)
        # 未安装时不得继续用假族名（否则整页方框，空格也变方框）
        self.assertIn(font.family().lower(), installed)
        self.assertNotIn("ZZNoSuchFontFamilyXYZ", font.family())

    def test_font_name_weight(self):
        self.assertTrue(_font_name_implies_bold("阿里巴巴普惠体 3.0 95 ExtraBold"))
        self.assertTrue(_font_name_implies_bold("阿里巴巴普惠体 3.0 105 Heavy"))
        self.assertFalse(_font_name_implies_bold("微软雅黑"))
        self.assertEqual(
            weight_from_font_name("阿里巴巴普惠体 3.0 95 ExtraBold"),
            QFont.Weight.ExtraBold,
        )

    def test_encode_parse_runs_keeps_fullwidth_space(self):
        raw = encode_rich_runs(
            [("价", "#FFFFFF"), ("\u3000\u3000", "#FF0000"), ("值", "#00FF00")]
        )
        self.assertTrue(raw.startswith(RICH_RUNS_MARK))
        runs = parse_rich_runs(raw)
        self.assertEqual(len(runs), 3)
        self.assertEqual(runs[1][0], "\u3000\u3000")
        self.assertEqual(plain_from_content(raw), "价\u3000\u3000值")

    def test_apply_runs_colors_and_spaces(self):
        style = ElementStyle(color="#FFFFFF", font_size=24)
        item = QGraphicsTextItem()
        item.setFont(QFont("Arial", 24))
        content = encode_rich_runs(
            [
                ("A", "#FF0000"),
                ("\u3000", "#00FF00"),
                ("B", "#0000FF"),
            ]
        )
        apply_text_content(item, content, style)
        self.assertEqual(item.toPlainText(), "A\u3000B")
        doc = item.document()
        colors = []
        for i in range(3):
            cur = QTextCursor(doc)
            cur.setPosition(i)
            cur.setPosition(i + 1, QTextCursor.MoveMode.KeepAnchor)
            colors.append(cur.charFormat().foreground().color().name().upper())
        self.assertEqual(colors[0], "#FF0000")
        self.assertEqual(colors[1], "#00FF00")
        self.assertEqual(colors[2], "#0000FF")

    def test_legacy_html_still_applies(self):
        style = ElementStyle(color="#FFFFFF", font_size=20)
        item = QGraphicsTextItem()
        item.setFont(QFont("Arial", 20))
        html = (
            RICH_HTML_MARK
            + '<span style="color:#AABBCC">你好</span>'
            + '<span style="color:#112233">&#12288;世界</span>'
        )
        apply_text_content(item, html, style)
        self.assertEqual(item.toPlainText(), "你好\u3000世界")
        self.assertEqual(plain_from_content(html), "你好\u3000世界")

    def test_sample_pptx_group_scale_and_defRPr(self):
        """组合子形状须乘 chExt 缩放；黄字常在 defRPr。"""
        pptx = Path.home() / "Desktop" / "重识生命的价值2.pptx"
        if not pptx.is_file():
            self.skipTest("sample pptx not on Desktop")
        from pptx import Presentation

        prs = Presentation(str(pptx))
        sx = 1920 / float(prs.slide_width)
        sy = 1080 / float(prs.slide_height)

        # 第 4 页副标题：defRPr 黄
        found_yellow = False
        for shape, *_rest in _iter_shapes(prs.slides[3].shapes):
            if not getattr(shape, "has_text_frame", False):
                continue
            if "基督如何恢复我们" not in (shape.text_frame.text or ""):
                continue
            _text, style = _shape_text_content(shape, sx=sx)
            self.assertEqual(style.color.upper(), "#F4D99B")
            found_yellow = True
        self.assertTrue(found_yellow)

        # 第 7 页标题：组合内不再是 8×8
        found_title = False
        for shape, _l, _t, w, h, *_rest in _iter_shapes(prs.slides[6].shapes):
            if not getattr(shape, "has_text_frame", False):
                continue
            if "你的生命" not in (shape.text_frame.text or ""):
                continue
            self.assertGreater(w * sx, 200)
            self.assertGreater(h * sy, 80)
            text, _style = _shape_text_content(shape, sx=sx)
            self.assertTrue(text.startswith(RICH_RUNS_MARK))
            self.assertIn("#F2D08A", text.upper())
            found_title = True
        self.assertTrue(found_title)


if __name__ == "__main__":
    unittest.main()
